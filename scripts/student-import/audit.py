"""Read-only workbook audit. No database connection or writes.
Usage: python scripts/student-import/audit.py /path/to/students.xlsx
Requires openpyxl. Prints aggregate counts and source cell references, not names/phones.
"""
import collections
import hashlib
import json
import re
import sys
from pathlib import Path

import openpyxl

def text(value):
    if value is None:
        return ""
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        return str(int(value)) if value == int(value) else str(value)
    return " ".join(str(value).split())

def audit(path):
    config = json.loads(Path(__file__).with_name("mapping.json").read_text(encoding="utf-8-sig"))
    workbook = openpyxl.load_workbook(path, data_only=False)
    report = {
        "sourceSha256": hashlib.sha256(path.read_bytes()).hexdigest(),
        "academicYear": config["academicYear"], "studentCount": 0,
        "sheets": [], "cohortCounts": collections.Counter(),
        "houseCounts": collections.Counter(), "studyStreamCounts": collections.Counter(),
        "issues": [], "phoneRowsWithMultipleContacts": [],
        "phoneRowsRequiringFormatReview": [],
        "duplicateNameCandidates": [], "duplicateNameParentCandidates": [],
    }
    names, identities = collections.defaultdict(list), collections.defaultdict(list)
    def issue(sheet, row, field, reason):
        report["issues"].append({"sheet": sheet, "row": row, "field": field, "reason": reason})
    for sheet in workbook:
        mapping = config["sheets"].get(sheet.title)
        if not mapping:
            issue(sheet.title, 1, "sheet", "Unmapped sheet")
            continue
        headers = [text(c.value) for c in sheet[1]]
        required = {"StudentID", "Name", "ParentName", "Mobile", "Class", "Group"}
        if "streams" in mapping:
            required.add("Stream")
        if len(headers) != len(set(headers)) or set(headers) != required:
            issue(sheet.title, 1, "headers", "Unexpected or duplicate headers")
            continue
        count = 0
        missing = collections.Counter()
        for row_number, values in enumerate(sheet.iter_rows(min_row=2, values_only=True), 2):
            if not any(text(v) for v in values):
                continue
            record = dict(zip(headers, values))
            count += 1
            ref = f"{sheet.title}!{row_number}"
            for field, value in record.items():
                if not text(value):
                    missing[field] += 1
                if isinstance(value, str) and value.startswith("="):
                    issue(sheet.title, row_number, field, "Formula cell: review source and cached value; never evaluate as code")
            name, parent = text(record["Name"]).casefold(), text(record["ParentName"]).casefold()
            if not name:
                issue(sheet.title, row_number, "Name", "Missing student name")
            else:
                names[name].append(ref)
                identities[(name, parent)].append(ref)
            class_value = text(record["Class"]).upper()
            if class_value != mapping["class"]:
                issue(sheet.title, row_number, "Class", "Blank class: proposed sheet-name fallback" if not class_value else "Class conflicts with sheet")
            stream = text(record.get("Stream")).upper()
            cohort = mapping.get("cohort") or mapping.get("cohortsByStream", {}).get(stream)
            if cohort:
                report["cohortCounts"][cohort] += 1
            else:
                issue(sheet.title, row_number, "Stream", "Cannot resolve cohort")
            if "streams" in mapping:
                normalized = mapping["streams"].get(stream)
                if normalized:
                    report["studyStreamCounts"][f"{cohort}:{normalized}"] += 1
                else:
                    issue(sheet.title, row_number, "Stream", "Unknown study stream")
            group = text(record["Group"]).upper()
            house = config["houseAliases"].get(group)
            report["houseCounts"][house or "UNASSIGNED"] += 1
            if group and not house:
                issue(sheet.title, row_number, "Group", "Unknown house")
            mobile = record["Mobile"]
            if text(mobile):
                # This is format triage, not international telephone-number validation.
                pieces = [v.strip() for v in re.split(r"[\n,;/]+", str(mobile)) if v.strip()]
                if len(pieces) > 1:
                    report["phoneRowsWithMultipleContacts"].append(ref)
                if isinstance(mobile, (int, float)):
                    pieces = [text(mobile)]
                simple = len(pieces) == 1 and re.fullmatch(r"[6-9]\d{9}", pieces[0])
                if not simple:
                    report["phoneRowsRequiringFormatReview"].append(ref)
        report["studentCount"] += count
        report["sheets"].append({"sheet": sheet.title, "studentCount": count, "headers": headers, "missing": dict(missing)})
    for missing_sheet in set(config["sheets"]) - set(workbook.sheetnames):
        issue(missing_sheet, 1, "sheet", "Expected sheet missing")
    report["duplicateNameCandidates"] = [refs for refs in names.values() if len(refs) > 1]
    report["duplicateNameParentCandidates"] = [refs for refs in identities.values() if len(refs) > 1]
    return report

if __name__ == "__main__":
    if len(sys.argv) != 2:
        raise SystemExit("Usage: python scripts/student-import/audit.py workbook.xlsx")
    print(json.dumps(audit(Path(sys.argv[1])), indent=2))
